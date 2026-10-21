#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
find_pattern_batch.py — 批量图案检索（语料索引一次，多查询复用）

与 find_pattern.py 的区别：本脚本面向"几百张查询 × 几百张语料"的批量场景。
语料特征（标记表+标定）只建一次索引，N 张查询不用把语料重复读 N 遍
（实测：198×244 若逐查询重建要 2 小时+，索引后只需几分钟）。

流程：
  1) 建索引：每张语料图 → 标记表 + 刻度标定（--index-json 可缓存复用）
  2) 粗筛：查询标记 × 语料标记 → 位移投票 → 候选峰（覆盖率 ≥ --cand-cov）
  3) 复核：只对够格的候选载入像素算残差 NCC（覆盖率+NCC 双证据）
  4) 输出：汇总 XLSX + CSV + 每查询命中复核图

CLI 用法：
  python find_pattern_batch.py --queries 查询目录 --corpus 语料目录 --out 输出目录
  python find_pattern_batch.py ... --limit 5 --tag pilot      # 试跑前 5 张

程序接口（GUI 复用）：run_batch()，进度/中止通过回调注入，见函数注释。

注意：本脚本固定 1.0 尺度（本数据集已确认查询图与语料 1:1）；
      被缩放的查询请用 find_pattern.py --scales auto。
"""
import argparse
import glob
import json
import os
import sys
import time

# --- 模块定位：开发态从同目录导入核心算法；单文件打包时这两行会被 build_single_ui.py 替换 ---
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import find_pattern as fp  # noqa: E402


# ---------------- 参数集合（CLI 与 GUI 共用） ----------------
class BatchParams:
    """一次批量检索的全部可调参数。默认值 = 0646 实战调优后的取值。"""

    def __init__(self, cand_cov=0.5, min_cov=0.75, min_ncc=0.4,
                 ncc_batch=96, ncc_max=400, stop_ncc=0.6,
                 strong_ncc=0.7, top=2):
        self.cand_cov = cand_cov      # 进入 NCC 复核的最低覆盖率
        self.min_cov = min_cov        # 判定命中的最低覆盖率
        self.min_ncc = min_ncc        # 判定命中的最低 NCC
        self.ncc_batch = ncc_batch    # 后续每批 NCC 候选数（首批固定 32）
        self.ncc_max = ncc_max        # 每查询 NCC 复核上限
        self.stop_ncc = stop_ncc      # 出现该强度的 NCC 即停止扩大复核
        self.strong_ncc = strong_ncc  # 主命中达到该值记『是』，否则记『疑似』
        self.top = top                # 每查询保存复核图的命中数


def _reporter(progress):
    """统一进度出口：有回调走回调（GUI），没有就打印（CLI）。"""
    if progress is not None:
        return progress

    def _p(_stage, _done, _total, msg):
        if msg:
            print(msg)
    return _p


# ---------------- 查询收集 ----------------
def collect_query_list(q, limit=None):
    """查询收集：目录递归 / 通配符 / 单文件。

    （区别于核心的 collect_queries：本函数可只取前 limit 张，供试跑用）
    """
    if os.path.isdir(q):
        out = []
        for root, dirs, names in os.walk(q):
            dirs[:] = [d for d in dirs if d != "__pycache__"]
            out += [os.path.join(root, n) for n in names
                    if n.lower().endswith(fp.IMG_EXTS)]
        out = sorted(out)
    else:
        out = sorted(p for p in glob.glob(q, recursive=True)
                     if p.lower().endswith(fp.IMG_EXTS))
    return out[:limit] if limit else out


# ---------------- 语料索引 ----------------
def build_index(files, cache_path, progress=None, should_stop=None):
    """语料索引：标记表 + 标定。缓存命中且集合/调色板一致则直接复用。

    返回索引列表；被中止（should_stop 变真）时返回 None（不写缓存）。
    """
    rep = _reporter(progress)
    if cache_path and os.path.exists(cache_path):
        try:
            with open(cache_path, encoding="utf-8") as f:
                idx = json.load(f)
            if ([os.path.abspath(x["path"]) for x in idx.get("files", [])]
                    == [os.path.abspath(x) for x in files]
                    and idx.get("palette") == json.dumps(fp.PALETTE)):
                rep("index", len(files), len(files),
                    f"[index] 复用缓存（{len(files)} 张）: {cache_path}")
                return idx["files"]
            rep("index", 0, len(files), "[index] 缓存与当前语料集合/调色板不一致，重建")
        except Exception as e:
            rep("index", 0, len(files), f"[index] 缓存不可用（{e}），重建")
    meta = []
    n = len(files)
    for i, f in enumerate(files):
        if should_stop is not None and should_stop():
            return None
        a = fp.load_rgb(f)
        marks = fp.extract_marks(a)
        pxm, x0 = fp.calibrate(a)
        meta.append({"path": os.path.abspath(f), "marks": marks,
                     "pxm": pxm, "x0": x0})
        msg = f"[index] {i + 1}/{n}" if ((i + 1) % 50 == 0 or i + 1 == n) else ""
        rep("index", i + 1, n, msg)
    if cache_path:
        with open(cache_path, "w", encoding="utf-8") as f:
            json.dump({"palette": json.dumps(fp.PALETTE), "files": meta},
                      f, ensure_ascii=False)
        rep("index", n, n, f"[index] 缓存写到 {cache_path}")
    return meta


# ---------------- XLSX 输出 ----------------
def write_xlsx(path, summary_rows, cand_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8F0E0")
    ok_fill = PatternFill("solid", fgColor="E1F5EE")
    sus_fill = PatternFill("solid", fgColor="FAEEDA")
    no_fill = PatternFill("solid", fgColor="F1EFE8")

    ws = wb.active
    ws.title = "检索结果"
    cols = ["查询图", "是否命中", "命中图片", "位移(px)", "覆盖率", "NCC",
            "区间(px)", "米区间(相对该图0点)", "次选图片", "复核图", "备注"]
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
    for r in summary_rows:
        ws.append([r.get(c, "") for c in cols])
    for i in range(2, ws.max_row + 1):
        v = ws.cell(row=i, column=2).value
        fill = ok_fill if v == "是" else (sus_fill if v == "疑似" else no_fill)
        for c in range(1, len(cols) + 1):
            ws.cell(row=i, column=c).fill = fill
    widths = [34, 8, 24, 10, 8, 7, 14, 18, 24, 40, 34]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("候选明细")
    cols2 = ["查询图", "候选图片", "位移(px)", "覆盖率", "NCC", "是否过门"]
    ws2.append(cols2)
    for c in range(1, len(cols2) + 1):
        cell = ws2.cell(row=1, column=c)
        cell.font = head_font
        cell.fill = head_fill
    for r in cand_rows:
        ws2.append([r.get(c, "") for c in cols2])
    for i, w in enumerate([34, 24, 10, 8, 7, 9], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    wb.save(path)


# ---------------- 单查询检索 ----------------
def search_one(qpath, q_rgb, Q, index, prm):
    """返回 (hits, accepted, n_ncc)。hits 已按 (覆盖率, NCC) 排序。"""
    hits = []
    for meta in index:
        hist = fp.vote_shifts(Q, meta["marks"], 1.0)
        if not hist:
            continue
        for p, _sc in fp.find_peaks(hist):
            cov, sup, miss = fp.coverage_at(Q, meta["marks"], p, 1.0)
            if cov >= prm.cand_cov:
                sd = (sum(fp.pair_shape_dist(q, c) for q, c in sup) / len(sup)
                      if sup else 9.0)
                hits.append({"file": meta["path"], "shift": float(p), "cov": cov,
                             "n_sup": len(sup), "total": len(sup) + len(miss),
                             "shape": sd, "ncc": None, "meta": meta})
    # 单标记查询会出现大量 cov 相同的并列候选（录里同形标记极多）：
    # 按 (覆盖率, 形状距离) 排序——真位置的标记形状必定最接近，先被 NCC 覆盖
    hits.sort(key=lambda h: (-h["cov"], h["shape"]))
    dedup = []
    for h in hits:
        if any(h["file"] == o["file"] and abs(h["shift"] - o["shift"]) <= 2
               for o in dedup):
            continue
        dedup.append(h)

    # NCC 复核（分批升级）：先查形状最接近的 32 个；出现 >= stop_ncc 的强结果
    # 立即停；否则按批扩大（上限 ncc_max）。单候选约 5ms 互相关 +
    # 约 60ms 载图/残差图（同文件缓存复用）。
    dq = fp.dmap(q_rgb, Q)
    dcache = {}
    checked = 0
    pos = 0
    while pos < len(dedup) and checked < prm.ncc_max:
        chunk = dedup[pos:pos + (32 if checked == 0 else prm.ncc_batch)]
        for h in chunk:
            dc = dcache.get(h["file"])
            if dc is None:
                dc = dcache[h["file"]] = fp.dmap(fp.load_rgb(h["file"]),
                                                 h["meta"]["marks"])
            h["ncc"] = fp.ncc_at_pos(dq, dc, h["shift"])
            checked += 1
        best_ok = max((h["ncc"] for h in dedup[:checked]
                       if h["ncc"] is not None), default=-1.0)
        if best_ok >= prm.stop_ncc:
            break
        pos += len(chunk)

    Wq = q_rgb.shape[1]
    for h in dedup:
        meta = h.pop("meta")
        P = h["shift"]
        h["px_range"] = [round(P), round(P + Wq)]
        if meta["pxm"] and meta["x0"] is not None:
            h["m_range"] = [round((P - meta["x0"]) / meta["pxm"], 3),
                            round((P + Wq - meta["x0"]) / meta["pxm"], 3)]
        else:
            h["m_range"] = None
    dedup.sort(key=lambda h: (-h["cov"], -(h["ncc"] or -1)))
    # 双证据齐备才算命中：NCC 没算出来的候选不放行（宁缺勿假）
    accepted = [h for h in dedup
                if h["cov"] >= prm.min_cov
                and h["ncc"] is not None and h["ncc"] >= prm.min_ncc]
    return dedup, accepted, checked


# ---------------- 主流程（CLI 与 GUI 共用） ----------------
def run_batch(queries, corpus, out_dir, prm=None, tag="", limit=None,
              index_json=None, progress=None, should_stop=None):
    """批量检索主流程。

    参数
    ------
    queries / corpus : 查询路径与语料路径（目录 / 通配符 / 单文件）
    out_dir          : 输出目录（复核图、结果表都写到这里）
    prm              : BatchParams；None 用默认值
    tag              : 输出文件名后缀
    limit            : 只跑前 N 张查询（试跑用），None 不限
    index_json       : 语料索引缓存路径（建议放稳定目录，跨批次复用）
    progress         : progress(stage, done, total, msg)
                       stage ∈ {"index","search","msg","error","done"}
                       GUI 用它刷进度条与日志；None 则把 msg 打印到 stdout
    should_stop      : should_stop() -> bool；返回真则中止（建索引每题、
                       每张查询前各检查一次），已完成部分照常写入结果

    返回统计 dict（含输出文件路径）。
    """
    prm = prm or BatchParams()
    rep = _reporter(progress)
    os.makedirs(out_dir, exist_ok=True)
    index_json = index_json or os.path.join(out_dir, "_corpus_index.json")

    qlist = collect_query_list(queries, limit)
    corpus_files = fp.collect_corpus(corpus, [out_dir])
    if not qlist or not corpus_files:
        msg = f"[err] 查询 {len(qlist)} 张 / 语料 {len(corpus_files)} 张，无法运行"
        rep("error", 0, 0, msg)
        return {"total": 0, "strong": 0, "suspect": 0, "miss": 0, "empty": 0,
                "stopped": False, "xlsx": "", "csv": "", "out": out_dir,
                "error": msg}
    rep("msg", 0, 0, f"batch v0.2 | 查询 {len(qlist)} 张 | 语料 {len(corpus_files)} 张")

    index = build_index(corpus_files, index_json,
                        progress=progress, should_stop=should_stop)
    if index is None:
        rep("msg", 0, 0, "[stop] 已在建索引阶段中止，未产生结果")
        return {"total": len(qlist), "strong": 0, "suspect": 0, "miss": 0,
                "empty": 0, "stopped": True, "xlsx": "", "csv": "",
                "out": out_dir, "error": ""}

    summary_rows, cand_rows = [], []
    n_strong = n_susp = n_miss = n_empty = 0
    stopped = False
    for qi, qp in enumerate(qlist, 1):
        if should_stop is not None and should_stop():
            stopped = True
            rep("msg", qi - 1, len(qlist),
                f"[stop] 用户中止：已完成 {qi - 1}/{len(qlist)} 张，已完成部分将写入结果")
            break
        qstem = os.path.splitext(os.path.basename(qp))[0]
        # 单张查询的任何异常都不允许中断全量：包一层 catch，记行继续
        try:
            q_rgb = fp.load_rgb(qp)
            Q = fp.extract_marks(q_rgb)
            if not Q:
                n_empty += 1
                summary_rows.append({"查询图": os.path.basename(qp), "是否命中": "否",
                                     "备注": "图中未提取到彩色标记", "查询路径": qp})
                cand_rows.append({"查询图": os.path.basename(qp), "候选图片": "—",
                                  "是否过门": "无标记"})
                rep("search", qi, len(qlist),
                    f"[{qi}/{len(qlist)}] [{time.strftime('%H:%M:%S')}] {qstem}: 无标记")
                continue

            hits, accepted, n_ncc = search_one(qp, q_rgb, Q, index, prm)

            odir = os.path.join(out_dir, "复核图", qstem)
            overlay_rel = ""
            if accepted:
                os.makedirs(odir, exist_ok=True)
                for i, h in enumerate(accepted[:prm.top], 1):
                    c_rgb = fp.load_rgb(h["file"])
                    png = os.path.join(odir, f"rank{i}_{os.path.splitext(os.path.basename(h['file']))[0]}"
                                             f"_x{int(round(h['shift']))}.png")
                    fullp = png.replace(".png", "_full.png")
                    fp.draw_overlay(q_rgb, Q, c_rgb,
                                    {"shift": h["shift"], "scale": 1.0}, png, fullp)
                    if i == 1:
                        overlay_rel = os.path.relpath(png, out_dir)

            if accepted:
                t = accepted[0]
                strong = t["ncc"] >= prm.strong_ncc
                if strong:
                    n_strong += 1
                else:
                    n_susp += 1
                # 次选：排除主文件；只取 NCC≥0.6 的有效条目；去重（同一位置常出现在相邻两张 Step）
                t2 = "、".join(list(dict.fromkeys(
                    os.path.basename(a["file"]) for a in accepted[1:6]
                    if a["file"] != t["file"] and a["ncc"] is not None
                    and a["ncc"] >= 0.6))[:2])
                row = {"查询图": os.path.basename(qp),
                       "是否命中": "是" if strong else "疑似",
                       "命中图片": os.path.basename(t["file"]),
                       "位移(px)": t["shift"], "覆盖率": round(t["cov"], 3),
                       "NCC": None if t["ncc"] is None else round(t["ncc"], 3),
                       "区间(px)": f"[{t['px_range'][0]},{t['px_range'][1]})",
                       "米区间(相对该图0点)": (f"{t['m_range'][0]:.2f}~{t['m_range'][1]:.2f}m"
                                              if t["m_range"] else ""),
                       "次选图片": t2,
                       "复核图": overlay_rel,
                       "备注": "" if strong else f"命中强度较弱（NCC<{prm.strong_ncc}），建议人工复核",
                       "查询路径": qp}
                rep("search", qi, len(qlist),
                    f"[{qi}/{len(qlist)}] [{time.strftime('%H:%M:%S')}] {qstem}: "
                    f"{'命中' if strong else '疑似'} {os.path.basename(t['file'])} "
                    f"shift={t['shift']:.0f} cov={t['cov']:.2f} ncc={t['ncc']:.3f}"
                    f"（NCC复核{n_ncc}个）")
            else:
                n_miss += 1
                best = hits[0] if hits else None
                note = (f"最高候选 {os.path.basename(best['file'])} cov={best['cov']:.2f}"
                        + (f" ncc={best['ncc']:.2f}" if best and best["ncc"] is not None else "")
                        + "（未过门）") if best else "无候选（覆盖率均<0.5）"
                row = {"查询图": os.path.basename(qp), "是否命中": "否",
                       "备注": note, "查询路径": qp}
                rep("search", qi, len(qlist),
                    f"[{qi}/{len(qlist)}] [{time.strftime('%H:%M:%S')}] {qstem}: "
                    f"未命中（{note}，NCC复核{n_ncc}个）")
            summary_rows.append(row)

            for h in hits[:10]:
                cand_rows.append({
                    "查询图": os.path.basename(qp),
                    "候选图片": os.path.basename(h["file"]),
                    "位移(px)": h["shift"],
                    "覆盖率": round(h["cov"], 3),
                    "NCC": None if h["ncc"] is None else round(h["ncc"], 3),
                    "是否过门": "是" if h in accepted else "否",
                })
        except Exception as e:
            import traceback
            traceback.print_exc()
            summary_rows.append({"查询图": os.path.basename(qp), "是否命中": "否",
                                 "备注": f"处理异常: {type(e).__name__}: {e}",
                                 "查询路径": qp})
            cand_rows.append({"查询图": os.path.basename(qp), "候选图片": "—",
                              "是否过门": "异常"})
            rep("search", qi, len(qlist),
                f"[{qi}/{len(qlist)}] [{time.strftime('%H:%M:%S')}] {qstem}: "
                f"异常 {type(e).__name__}: {e}")

    tag_s = f"_{tag}" if tag else ""
    xlsx = os.path.join(out_dir, f"图案检索结果{tag_s}.xlsx")
    write_xlsx(xlsx, summary_rows, cand_rows)
    csvp = os.path.join(out_dir, f"图案检索结果{tag_s}.csv")
    fp.write_csv(csvp, summary_rows,
                 ["查询图", "是否命中", "命中图片", "位移(px)", "覆盖率", "NCC",
                  "区间(px)", "米区间(相对该图0点)", "次选图片", "复核图", "备注", "查询路径"])
    rep("done", len(qlist), len(qlist),
        f"\n== 完成：命中(强) {n_strong} / 疑似 {n_susp} / 未命中 {n_miss} / 无标记 {n_empty}"
        f"  （共 {len(qlist)}）" + ("  [已中止]" if stopped else ""))
    rep("done", len(qlist), len(qlist), f"[out] {xlsx}")
    rep("done", len(qlist), len(qlist), f"[out] {csvp}")
    return {"total": len(qlist), "strong": n_strong, "suspect": n_susp,
            "miss": n_miss, "empty": n_empty, "stopped": stopped,
            "xlsx": xlsx, "csv": csvp, "out": out_dir, "error": ""}


# ---------------- CLI ----------------
def main():
    ap = argparse.ArgumentParser(description="批量图案检索（语料索引复用）")
    ap.add_argument("--queries", required=True, help="查询图目录 / 通配符 / 单文件")
    ap.add_argument("--corpus", required=True, help="语料目录（递归）")
    ap.add_argument("--out", required=True, help="输出目录")
    ap.add_argument("--limit", type=int, default=None, help="只跑前 N 张查询（试跑用）")
    ap.add_argument("--tag", default="", help="输出文件名后缀，如 pilot")
    ap.add_argument("--cand-cov", type=float, default=0.5, help="进入 NCC 复核的最低覆盖率")
    ap.add_argument("--min-cov", type=float, default=0.75, help="判定命中的最低覆盖率")
    ap.add_argument("--min-ncc", type=float, default=0.4, help="判定命中的最低 NCC")
    ap.add_argument("--ncc-batch", type=int, default=96, help="后续每批 NCC 候选数（首批固定 32）")
    ap.add_argument("--ncc-max", type=int, default=400, help="每查询 NCC 复核上限")
    ap.add_argument("--stop-ncc", type=float, default=0.6, help="出现该强度的 NCC 即停止扩大复核")
    ap.add_argument("--strong-ncc", type=float, default=0.7,
                    help="主命中 NCC 达到该值记『是』，低于记『疑似』（低于 min-ncc 不记）")
    ap.add_argument("--top", type=int, default=2, help="每查询保存复核图的命中数")
    ap.add_argument("--index-json", default=None, help="语料索引缓存路径（跨批次复用）")
    args = ap.parse_args()

    prm = BatchParams(cand_cov=args.cand_cov, min_cov=args.min_cov,
                      min_ncc=args.min_ncc, ncc_batch=args.ncc_batch,
                      ncc_max=args.ncc_max, stop_ncc=args.stop_ncc,
                      strong_ncc=args.strong_ncc, top=args.top)
    stats = run_batch(args.queries, args.corpus, args.out, prm=prm,
                      tag=args.tag, limit=args.limit, index_json=args.index_json)
    return 2 if stats.get("error") else 0


if __name__ == "__main__":
    sys.exit(main())
